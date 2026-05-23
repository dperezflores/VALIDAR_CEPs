import time
import random
import os
import glob
import pandas as pd
import fitz

from docx import Document
from docx.shared import Inches

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import Select

from config import DOWNLOAD_DIR, URL_BANXICO_CEP
from utils import homologar_banco


class ValidadorBanxico:
    def __init__(self):
        self.doc = Document()
        self.ruta_word = os.path.join(DOWNLOAD_DIR, "Capturas_CEP.docx")

        self.driver = self._configurar_navegador()
        self.wait = WebDriverWait(self.driver, 120)

    def _configurar_navegador(self):
        options = Options()

        options.add_argument("--headless=new")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-gpu")
        options.add_argument("--window-size=1920,1080")

        # Chromium en Streamlit Cloud
        options.binary_location = "/usr/bin/chromium"

        prefs = {
            "download.default_directory": DOWNLOAD_DIR,
            "download.prompt_for_download": False,
            "plugins.always_open_pdf_externally": True
        }

        options.add_experimental_option("prefs", prefs)

        service = Service("/usr/bin/chromedriver")

        driver = webdriver.Chrome(
            service=service,
            options=options
        )

        return driver

    def _escribir_como_humano(self, elemento, texto: str):
        """Simula tecleo humano."""
        for letra in texto:
            elemento.send_keys(letra)
            time.sleep(random.uniform(0.05, 0.25))

    def procesar_registro(self, row: pd.Series) -> str:

        if (
            pd.isna(row["Clave de rastreo"])
            or str(row["Clave de rastreo"]).strip().upper() == "N/A"
        ):
            return "Omitido: Sin Clave de rastreo"

        try:
            self.driver.delete_all_cookies()
            self.driver.get(URL_BANXICO_CEP)

            time.sleep(random.uniform(3.0, 5.0))

            # =========================
            # 1. FECHA
            # =========================
            input_fecha = self.wait.until(
                EC.presence_of_element_located((By.ID, "input_fecha"))
            )

            fecha_raw = str(row["Fecha de pago"]).strip().lower()

            meses = {
                "ene": "01",
                "feb": "02",
                "mar": "03",
                "abr": "04",
                "may": "05",
                "jun": "06",
                "jul": "07",
                "ago": "08",
                "sep": "09",
                "oct": "10",
                "nov": "11",
                "dic": "12",
            }

            for letras, numero in meses.items():
                if letras in fecha_raw:
                    fecha_raw = fecha_raw.replace(letras, numero)

            try:
                fecha_limpia = pd.to_datetime(fecha_raw).strftime("%d-%m-%Y")
            except:
                fecha_limpia = fecha_raw

            script_js = """
                var input = arguments[0];
                input.value = arguments[1];
                input.dispatchEvent(new Event('input', { bubbles: true }));
                input.dispatchEvent(new Event('change', { bubbles: true }));
            """

            self.driver.execute_script(
                script_js,
                input_fecha,
                fecha_limpia
            )

            time.sleep(random.uniform(0.5, 1.5))

            # =========================
            # 2. CRITERIO
            # =========================
            Select(
                self.driver.find_element(
                    By.ID,
                    "input_tipoCriterio"
                )
            ).select_by_visible_text("Clave de rastreo")

            time.sleep(random.uniform(0.5, 1.0))

            input_criterio = self.driver.find_element(
                By.ID,
                "input_criterio"
            )

            self._escribir_como_humano(
                input_criterio,
                str(row["Clave de rastreo"])
            )

            time.sleep(random.uniform(0.5, 1.0))

            # =========================
            # 3. BANCOS
            # =========================
            banco_emisor = homologar_banco(
                row["Institución emisora"]
            )

            banco_receptor = homologar_banco(
                row["Institución receptora"]
            )

            Select(
                self.driver.find_element(By.ID, "input_emisor")
            ).select_by_visible_text(banco_emisor)

            time.sleep(random.uniform(0.5, 1.0))

            Select(
                self.driver.find_element(By.ID, "input_receptor")
            ).select_by_visible_text(banco_receptor)

            time.sleep(random.uniform(0.5, 1.0))

            # =========================
            # 4. CUENTA Y MONTO
            # =========================
            input_cuenta = self.driver.find_element(
                By.ID,
                "input_cuenta"
            )

            self._escribir_como_humano(
                input_cuenta,
                str(row["Cuenta beneficiaria"])
            )

            time.sleep(random.uniform(0.5, 1.0))

            monto_limpio = (
                str(row["Importe"])
                .replace(",", "")
                .replace("$", "")
                .strip()
            )

            input_monto = self.driver.find_element(
                By.ID,
                "input_monto"
            )

            self._escribir_como_humano(
                input_monto,
                monto_limpio
            )

            time.sleep(random.uniform(1.0, 2.0))

            # =========================
            # Registrar PDFs existentes
            # =========================
            pdfs_antes = set(
                glob.glob(
                    os.path.join(DOWNLOAD_DIR, "*.pdf")
                )
            )

            # =========================
            # 5. BOTÓN CONSULTAR
            # =========================
            try:
                btn_principal = self.driver.find_element(
                    By.ID,
                    "btn_Descargar"
                )
            except:
                btn_principal = self.driver.find_element(
                    By.ID,
                    "btn_Consultar"
                )

            self.driver.execute_script(
                "arguments[0].click();",
                btn_principal
            )

            # =========================
            # 6. Esperar modal PDF
            # =========================
            try:
                btn_pdf_modal = self.wait.until(
                    EC.element_to_be_clickable(
                        (
                            By.CSS_SELECTOR,
                            ".boton.boton-descarga-pdf"
                        )
                    )
                )

                time.sleep(random.uniform(1.0, 2.0))

                self.driver.execute_script(
                    "arguments[0].click();",
                    btn_pdf_modal
                )

            except:
                return "Rechazado por Banxico / Timeout CAPTCHA"

            # =========================
            # 7. Detectar nuevo PDF
            # =========================
            pdf_nuevo = None

            for _ in range(15):
                pdfs_despues = set(
                    glob.glob(
                        os.path.join(DOWNLOAD_DIR, "*.pdf")
                    )
                )

                archivos_nuevos = pdfs_despues - pdfs_antes

                if archivos_nuevos:
                    pdf_nuevo = archivos_nuevos.pop()
                    break

                time.sleep(1)

            if not pdf_nuevo:
                return "Error: Falla en descarga"

            # =========================
            # Convertir PDF a imagen
            # =========================
            documento_pdf = fitz.open(pdf_nuevo)
            pagina = documento_pdf.load_page(0)

            matriz_zoom = fitz.Matrix(2.0, 2.0)

            mapa_bits = pagina.get_pixmap(
                matrix=matriz_zoom
            )

            ruta_imagen_temp = "temp_pdf_render.png"

            mapa_bits.save(ruta_imagen_temp)

            documento_pdf.close()

            titulo = (
                f"{row.get('Número', '')} - "
                f"{row.get('Archivo Origen', '')}"
            )

            self.doc.add_heading(titulo, level=1)

            self.doc.add_picture(
                ruta_imagen_temp,
                width=Inches(6.0)
            )

            if os.path.exists(ruta_imagen_temp):
                os.remove(ruta_imagen_temp)

            return "Éxito (PDF descargado)"

        except Exception as e:
            return f"Error técnico: {type(e).__name__}"

    def procesar_lote(
        self,
        df: pd.DataFrame,
        funcion_progreso=None
    ) -> pd.DataFrame:

        resultados = []
        total = len(df)

        for index, row in df.iterrows():

            estado = self.procesar_registro(row)

            resultados.append({
                "Clave de rastreo": row.get(
                    "Clave de rastreo",
                    "N/A"
                ),
                "Estado": estado
            })

            if funcion_progreso:
                funcion_progreso(index + 1, total)

            if index < total - 1:
                time.sleep(
                    random.uniform(10.0, 15.0)
                )

        return pd.DataFrame(resultados)

    def cerrar(self):

        if len(self.doc.paragraphs) > 0:
            self.doc.save(self.ruta_word)

        if self.driver:
            self.driver.quit()
